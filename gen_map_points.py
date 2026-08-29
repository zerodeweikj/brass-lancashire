# -*- coding: utf-8 -*-
"""从 兰开夏建筑城市坐标.xlsx 解析像素坐标，生成 web/public/data/map_points.json。
表结构（纵向）：
  行1 = 城市id(英文)，行2 = 城市中文名，行3 = 城市像素坐标，行4 = 建筑槽位(类型+坐标)
  每列一座城市（B..Z）。
输出格式与 MapCalibrateScene._export 同构：
  { "map": "main_map.jpg", "locations": { id: {point:[x,y]} | {slots:[{types,x,y}]} } }
  额外附带 "points" 字段（每城单点）以兼容 MapScene 当前读取逻辑。
"""
import openpyxl, re, json, os

SRC = r"C:/Users/bwf/Desktop/兰开夏建筑城市坐标.xlsx"
OUT = r"D:/zhuoyouyizhi/lancashire/web/public/data/map_points.json"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Sheet1"]

slot_re = re.compile(r'([【】一-鿿]+?)\s*[\(（]\s*(\d+)\s*[\,，]\s*(\d+)\s*[\)）]')
coord_re = re.compile(r'(\d+)\s*[\,，]\s*(\d+)')

def parse_type(t: str):
    t = t.strip('【】')
    if '棉' in t and '煤' in t:
        return ["cotton", "coal"]
    if '棉花厂或港口' in t or '棉纺厂或港口' in t:
        return ["cotton", "port"]
    if '铁' in t:
        return ["iron"]
    if '棉花' in t or '棉纺' in t:
        return ["cotton"]
    if '煤' in t:
        return ["coal"]
    if '港口' in t:
        return ["port"]
    if '造船' in t:
        return ["shipyard"]
    return [t.strip()]

locations = {}
points = {}
n_slot = 0
for col in range(2, ws.max_column + 1):
    city_id = (ws.cell(row=1, column=col).value or "")
    if city_id is None:
        continue
    city_id = str(city_id).strip()
    if not city_id:
        continue
    name = str(ws.cell(row=2, column=col).value or "").strip()
    city_coord = str(ws.cell(row=3, column=col).value or "")
    slots_raw = str(ws.cell(row=4, column=col).value or "")
    m = coord_re.search(city_coord)
    city_pt = [int(m.group(1)), int(m.group(2))] if m else None

    slots = []
    if slots_raw and "无" not in slots_raw:
        for sm in slot_re.finditer(slots_raw):
            slots.append({
                "types": parse_type(sm.group(1)),
                "x": int(sm.group(2)),
                "y": int(sm.group(3)),
            })

    if not slots:
        locations[city_id] = {"point": city_pt}
    else:
        locations[city_id] = {"slots": slots}
        n_slot += len(slots)
    # points：每城单点（地图场景画城市标记/连线用），统一取城市中心坐标
    points[city_id] = city_pt
    n_slot_city = len(slots)
    print(f"{city_id:24s} {name:8s} city_pt={city_pt} slots={n_slot_city}")
    for s in slots:
        print(f"      {s['types']} -> ({s['x']},{s['y']})")

out = {"map": "main_map.jpg", "locations": locations, "points": points}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print(f"\n=> wrote {OUT}")
print(f"cities={len(locations)} slots_total={n_slot} city_points={len(locations)-sum(1 for v in locations.values() if 'slots' in v)}")
